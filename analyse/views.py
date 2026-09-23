import io
import re
import unicodedata
from datetime import datetime
from decimal import Decimal

from django.contrib import messages
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import render

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .forms import ImportExcelForm
from .services.importer import importer_fichier_excel

from django.db.models import Avg, Max, Min, Count, Q
from .models import (
    AnneeScolaire,
    Classe,
    Eleve,
    Resultat,
)




def _normaliser(texte):
    """Retire les accents et met en minuscules pour faciliter les
    comparaisons de niveaux ("2ndeA", "2NDE A", "2nde-A", ...)."""

    texte = texte or ""
    texte = unicodedata.normalize("NFKD", texte)
    texte = texte.encode("ascii", "ignore").decode("ascii")
    return texte.lower().strip()


def _info_tri_niveau(niveau):
    """Détermine à quel cycle appartient un niveau et calcule une clé
    de tri, à partir de la valeur brute stockée dans Classe.niveau."""

    n = _normaliser(niveau)

    # --- Collège : 6e, 5e, 4e, 3e (ou "6ème", "6 eme", ...) ---
    college = re.match(r"^([3-6])\s*e", n)

    if college:
        chiffre = int(college.group(1))
        # Ordre décroissant : 6e, 5e, 4e, 3e
        return "1er", (0, -chiffre, "")

    # --- Lycée : Seconde / Première / Terminale ---
    prefixes_lycee = [
        ("seconde", 0),
        ("2nde", 0),
        ("2de", 0),
        ("2e", 0),
        ("premiere", 1),
        ("1ere", 1),
        ("1re", 1),
        ("1e", 1),
        ("terminale", 2),
        ("tle", 2),
        ("term", 2),
    ]

    for prefixe, rang in prefixes_lycee:
        if n.startswith(prefixe):
            section = n[len(prefixe):].strip().upper()
            return "2nd", (1, rang, section)

    # --- Niveau non reconnu : placé en fin de tableau ---
    return "2nd", (2, 0, n)


CHAMPS_NUMERIQUES_NIVEAU = (
    "nombre_classes",
    "classe_f", "classe_g", "classe_t",
    "non_classe_f", "non_classe_g", "non_classe_t",
    "sup10_f", "sup10_g", "sup10_t",
    "moy_f", "moy_g", "moy_t",
    "inf_f", "inf_g", "inf_t",
)


def _nouvelle_ligne_niveau(niveau):

    ligne = {"niveau": niveau, "is_total": False, "classes_ids": set()}

    for champ in CHAMPS_NUMERIQUES_NIVEAU:
        ligne[champ] = 0

    ligne["somme_moyennes"] = Decimal("0")
    ligne["nb_moyennes"] = 0

    return ligne


def _finaliser_ligne(ligne):
    """Calcule nombre_classes, pourcentages et moyenne finale d'une ligne."""

    if "classes_ids" in ligne:
        ligne["nombre_classes"] = len(ligne.pop("classes_ids"))

    total_classe = ligne["classe_t"] or 0

    ligne["sup10_pct"] = (
        round(ligne["sup10_t"] / total_classe * 100, 2)
        if total_classe else 0
    )
    ligne["moy_pct"] = (
        round(ligne["moy_t"] / total_classe * 100, 2)
        if total_classe else 0
    )
    ligne["inf_pct"] = (
        round(ligne["inf_t"] / total_classe * 100, 2)
        if total_classe else 0
    )

    ligne["moyenne_niveau"] = (
        round(float(ligne["somme_moyennes"] / ligne["nb_moyennes"]), 2)
        if ligne["nb_moyennes"] else None
    )

    return ligne


def _construire_lignes_par_niveau(resultats):
    """Parcourt les résultats filtrés et construit une ligne de
    statistiques par niveau de classe."""

    lignes_par_niveau = {}

    for resultat in resultats:

        niveau = resultat.classe.niveau or "Non renseigné"

        ligne = lignes_par_niveau.setdefault(
            niveau, _nouvelle_ligne_niveau(niveau)
        )

        ligne["classes_ids"].add(resultat.classe_id)

        genre = resultat.eleve.genre
        moyenne = resultat.moyenne_trimestrielle

        est_fille = genre == "F"
        est_garcon = genre == "M"

        if moyenne is None or moyenne == 0:

            ligne["non_classe_t"] += 1

            if est_fille:
                ligne["non_classe_f"] += 1
            elif est_garcon:
                ligne["non_classe_g"] += 1

        else:

            ligne["classe_t"] += 1

            if est_fille:
                ligne["classe_f"] += 1
            elif est_garcon:
                ligne["classe_g"] += 1

            ligne["somme_moyennes"] += moyenne
            ligne["nb_moyennes"] += 1

            if moyenne >= 10:
                ligne["sup10_t"] += 1
                if est_fille:
                    ligne["sup10_f"] += 1
                elif est_garcon:
                    ligne["sup10_g"] += 1

            elif moyenne >= Decimal("8.5"):
                ligne["moy_t"] += 1
                if est_fille:
                    ligne["moy_f"] += 1
                elif est_garcon:
                    ligne["moy_g"] += 1

            else:
                ligne["inf_t"] += 1
                if est_fille:
                    ligne["inf_f"] += 1
                elif est_garcon:
                    ligne["inf_g"] += 1

    return [_finaliser_ligne(l) for l in lignes_par_niveau.values()]


def _fusionner_lignes(lignes, libelle):
    """Additionne plusieurs lignes de niveau pour obtenir une ligne
    de total (1er cycle / 2nd cycle / général)."""

    fusion = _nouvelle_ligne_niveau(libelle)
    fusion.pop("classes_ids")
    fusion["nombre_classes"] = 0
    fusion["is_total"] = True

    for ligne in lignes:

        for champ in CHAMPS_NUMERIQUES_NIVEAU:
            fusion[champ] += ligne[champ]

        fusion["somme_moyennes"] += ligne["somme_moyennes"]
        fusion["nb_moyennes"] += ligne["nb_moyennes"]

    return _finaliser_ligne(fusion)


def calculer_tableau_statistiques_niveaux(resultats):
    """Construit la liste finale de lignes (dans l'ordre d'affichage)
    à utiliser directement dans le template : niveaux de 1er cycle,
    total 1er cycle, niveaux de 2nd cycle, total 2nd cycle, total
    général."""

    lignes = _construire_lignes_par_niveau(resultats)

    for ligne in lignes:
        cycle, cle_tri = _info_tri_niveau(ligne["niveau"])
        ligne["_cycle"] = cycle
        ligne["_cle_tri"] = cle_tri

    lignes_1er_cycle = sorted(
        (l for l in lignes if l["_cycle"] == "1er"),
        key=lambda l: l["_cle_tri"],
    )

    lignes_2nd_cycle = sorted(
        (l for l in lignes if l["_cycle"] == "2nd"),
        key=lambda l: l["_cle_tri"],
    )

    tableau = list(lignes_1er_cycle)

    if lignes_1er_cycle:
        tableau.append(
            _fusionner_lignes(lignes_1er_cycle, "Total 1er cycle")
        )

    tableau += lignes_2nd_cycle

    if lignes_2nd_cycle:
        tableau.append(
            _fusionner_lignes(lignes_2nd_cycle, "Total 2nd cycle")
        )

    if lignes_1er_cycle or lignes_2nd_cycle:
        tableau.append(
            _fusionner_lignes(lignes, "Total général")
        )

    return tableau


# =====================================================================
# TABLEAU EFFECTIFS / AFFECTÉS / REDOUBLANTS PAR NIVEAU (tableau "STAT_2")
# =====================================================================
#
# Hypothèses retenues (à ajuster si besoin) :
#   - "EFFECTIF"    = nombre d'élèves distincts (par sexe) ayant un
#                      résultat correspondant aux filtres appliqués,
#                      pour le niveau concerné.
#   - "AFFECTÉS"    = sous-ensemble de ces élèves dont le champ
#                      Eleve.statut contient "affecté" (comparaison
#                      insensible à la casse et aux accents).
#   - "REDOUBLANTS" = sous-ensemble de ces élèves dont
#                      Eleve.redoublant vaut True.
#   - "G" = genre 'M' (Garçon), "F" = genre 'F' (Fille).
#   - En plus du total 1er cycle / 2nd cycle, le 2nd cycle est
#     également sous-totalisé par palier (2nde, 1ère, Tle), comme
#     dans le modèle fourni.


CHAMPS_NUMERIQUES_EFFECTIFS = (
    "effectif_g", "effectif_f", "effectif_t",
    "affectes_g", "affectes_f", "affectes_t",
    "redoublants_g", "redoublants_f", "redoublants_t",
)


def _est_affecte(statut):
    return "affecte" in _normaliser(statut or "")


def _calculer_ligne_effectifs(niveau, eleves):

    ligne = {"niveau": niveau, "is_total": False}

    for champ in CHAMPS_NUMERIQUES_EFFECTIFS:
        ligne[champ] = 0

    for eleve in eleves:

        genre = eleve.genre
        est_garcon = genre == "M"
        est_fille = genre == "F"

        ligne["effectif_t"] += 1
        if est_garcon:
            ligne["effectif_g"] += 1
        elif est_fille:
            ligne["effectif_f"] += 1

        if _est_affecte(eleve.statut):
            ligne["affectes_t"] += 1
            if est_garcon:
                ligne["affectes_g"] += 1
            elif est_fille:
                ligne["affectes_f"] += 1

        if eleve.redoublant:
            ligne["redoublants_t"] += 1
            if est_garcon:
                ligne["redoublants_g"] += 1
            elif est_fille:
                ligne["redoublants_f"] += 1

    return ligne


def _construire_lignes_effectifs(resultats):
    """Construit, pour chaque niveau, l'ensemble des élèves distincts
    (par élève, pas par résultat) correspondant aux filtres, puis
    calcule les compteurs effectif / affectés / redoublants."""

    niveau_eleves = {}

    for resultat in resultats:

        niveau = resultat.classe.niveau or "Non renseigné"

        eleves_niveau = niveau_eleves.setdefault(niveau, {})
        eleves_niveau[resultat.eleve_id] = resultat.eleve

    return [
        _calculer_ligne_effectifs(niveau, eleves.values())
        for niveau, eleves in niveau_eleves.items()
    ]


def _fusionner_lignes_effectifs(lignes, libelle):

    fusion = {"niveau": libelle, "is_total": True}

    for champ in CHAMPS_NUMERIQUES_EFFECTIFS:
        fusion[champ] = sum(l[champ] for l in lignes)

    return fusion


NOMS_PALIERS_LYCEE = {
    0: "Total 2nde",
    1: "Total 1ere",
    2: "Total Tle",
}


def calculer_tableau_effectifs_niveaux(resultats):
    """Construit la liste finale de lignes (dans l'ordre d'affichage)
    pour le tableau effectifs / affectés / redoublants : niveaux de
    1er cycle, total 1er cycle, puis pour le 2nd cycle un
    sous-total par palier (2nde / 1ère / Tle), un total 2nd cycle,
    et enfin le total général."""

    lignes = _construire_lignes_effectifs(resultats)

    for ligne in lignes:
        cycle, cle_tri = _info_tri_niveau(ligne["niveau"])
        ligne["_cycle"] = cycle
        ligne["_cle_tri"] = cle_tri

    lignes_1er_cycle = sorted(
        (l for l in lignes if l["_cycle"] == "1er"),
        key=lambda l: l["_cle_tri"],
    )

    lignes_2nd_cycle = sorted(
        (l for l in lignes if l["_cycle"] == "2nd"),
        key=lambda l: l["_cle_tri"],
    )

    tableau = list(lignes_1er_cycle)

    if lignes_1er_cycle:
        tableau.append(
            _fusionner_lignes_effectifs(lignes_1er_cycle, "Total 1er Cycle")
        )

    paliers = {}
    niveaux_non_reconnus = []

    for ligne in lignes_2nd_cycle:
        if ligne["_cle_tri"][0] == 1:
            palier = ligne["_cle_tri"][1]
            paliers.setdefault(palier, []).append(ligne)
        else:
            niveaux_non_reconnus.append(ligne)

    for palier in sorted(paliers.keys()):

        lignes_palier = paliers[palier]
        tableau += lignes_palier

        tableau.append(
            _fusionner_lignes_effectifs(
                lignes_palier,
                NOMS_PALIERS_LYCEE.get(palier, "Total palier")
            )
        )

    tableau += niveaux_non_reconnus

    if lignes_2nd_cycle:
        tableau.append(
            _fusionner_lignes_effectifs(lignes_2nd_cycle, "Total 2e Cycle")
        )

    if lignes_1er_cycle or lignes_2nd_cycle:
        tableau.append(
            _fusionner_lignes_effectifs(lignes, "Total Général")
        )

    return tableau



# =====================================================================
# STATISTIQUES PAR MATIÈRE
# =====================================================================

SEUIL_REUSSITE = 10

MATIERES = {
    "Composition française": "composition_francaise",
    "Orthographe / Grammaire": "orthographe_grammaire",
    "Expression orale": "expression_orale",
    "Histoire-Géographie": "hist_geo",
    "Espagnol": "espagnol",
    "Allemand": "allemand",
    "Anglais": "anglais",
    "Français": "francais",
    "Conduite": "conduite",
    "EDHC": "edhc",
    "EPS": "eps",
    "SVT": "svt",
    "Physiques-Chimie": "physiques_chimie",
    "Mathématiques": "mathematiques",
    "Philosophie": "philosophie",
}


def calculer_statistiques_matieres(resultats):
    """Moyenne, effectif et taux de réussite (note >= 10) de chaque
    matière. Une seule requête SQL pour toutes les matières."""

    agregats = {}

    for champ in MATIERES.values():
        agregats[f"moy_{champ}"] = Avg(champ)
        agregats[f"nb_{champ}"] = Count(champ)
        agregats[f"ok_{champ}"] = Count(
            champ,
            filter=Q(**{f"{champ}__gte": SEUIL_REUSSITE})
        )

    totaux = resultats.aggregate(**agregats)

    statistiques = []

    for nom, champ in MATIERES.items():

        moyenne = totaux[f"moy_{champ}"]
        nombre = totaux[f"nb_{champ}"]
        reussis = totaux[f"ok_{champ}"]

        if moyenne is not None and nombre:
            moyenne = round(float(moyenne), 2)
            taux_reussite = round(reussis / nombre * 100, 2)
        else:
            moyenne = None
            taux_reussite = 0

        statistiques.append({
            "nom": nom,
            "champ": champ,
            "moyenne": moyenne,
            "nombre": nombre,
            "taux_reussite": taux_reussite,
        })

    return statistiques


# =====================================================================
# DONNÉES DES GRAPHIQUES DU DASHBOARD
# =====================================================================
#
# Trois graphiques sont affichés : répartition des élèves par tranche
# de moyenne, moyenne par niveau, et filles / garçons par tranche.
# Ils réutilisent le tableau "Statistiques par niveau" déjà calculé
# (aucune requête SQL supplémentaire). Les données sont transmises au
# template dans un dictionnaire sérialisé en JSON (filtre
# `json_script`), que Chart.js lit côté navigateur.


def construire_donnees_graphiques(tableau_niveaux):

    lignes_niveaux = [l for l in tableau_niveaux if not l["is_total"]]

    total = next(
        (l for l in tableau_niveaux if l["niveau"] == "Total général"),
        None
    )

    if not lignes_niveaux or total is None:
        return {"a_des_donnees": False}

    labels_tranches = [
        "≥ 10",
        "8,5 à < 10",
        "< 8,5",
        "Non classés",
    ]

    # ---- Répartition des élèves par tranche de moyenne ----

    tranches = {
        "labels": labels_tranches,
        "valeurs": [
            total["sup10_t"],
            total["moy_t"],
            total["inf_t"],
            total["non_classe_t"],
        ],
        "taux_reussite": (
            total["sup10_pct"] if total["classe_t"] else None
        ),
    }

    # ---- Moyenne par niveau ----

    niveaux = {
        "labels": [l["niveau"] for l in lignes_niveaux],
        "moyennes": [l["moyenne_niveau"] for l in lignes_niveaux],
    }

    # ---- Filles / garçons par tranche de moyenne ----

    genre = {
        "labels": labels_tranches,
        "filles": [
            total["sup10_f"], total["moy_f"],
            total["inf_f"], total["non_classe_f"],
        ],
        "garcons": [
            total["sup10_g"], total["moy_g"],
            total["inf_g"], total["non_classe_g"],
        ],
    }

    return {
        "a_des_donnees": True,
        "seuil": SEUIL_REUSSITE,
        "tranches": tranches,
        "niveaux": niveaux,
        "genre": genre,
    }


def dashboard(request):

    # -----------------------------
    # IMPORTATION EXCEL
    # -----------------------------

    if request.method == "POST":

        form = ImportExcelForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():

            fichier = form.cleaned_data["fichier"]

            try:

                resultat_import = importer_fichier_excel(
                    fichier
                )

                messages.success(
                    request,
                    f"Importation réussie ! "
                    f"{resultat_import['eleves']} élèves et "
                    f"{resultat_import['resultats']} résultats "
                    f"ont été importés."
                )

            except Exception as e:

                messages.error(
                    request,
                    f"Échec de l'importation : {str(e)}"
                )

        else:

            messages.error(
                request,
                "Échec de l'importation. "
                "Veuillez sélectionner un fichier Excel valide."
            )

    else:

        form = ImportExcelForm()


    # -----------------------------
    # FILTRES DU DASHBOARD
    # -----------------------------

    annee_id = request.GET.get("annee")
    classe_id = request.GET.get("classe")
    trimestre = request.GET.get("trimestre")
    niveau = request.GET.get("niveau")
    recherche = request.GET.get("recherche", "").strip()

    annees = AnneeScolaire.objects.all().order_by("-nom")
    classes = Classe.objects.all().order_by("nom")
    
    # Récupération des niveaux uniques existants dans la base de données
    niveaux_uniques = Classe.objects.values_list('niveau', flat=True).distinct().exclude(niveau='').order_by('niveau')

    resultats = Resultat.objects.select_related(
        "eleve",
        "annee_scolaire",
        "classe"
    )

    # Filtre par niveau
    if niveau:
        resultats = resultats.filter(
            classe__niveau=niveau
        )

    if annee_id:
        resultats = resultats.filter(
            annee_scolaire_id=annee_id
        )

    if classe_id:
        resultats = resultats.filter(
            classe_id=classe_id
        )

    if trimestre:
        resultats = resultats.filter(
            trimestre=trimestre
        )

    # Recherche par nom, prénom ou matricule
    if recherche:
        resultats = resultats.filter(
            Q(eleve__nom__icontains=recherche)
            | Q(eleve__prenoms__icontains=recherche)
            | Q(eleve__matricule__icontains=recherche)
        )

    # Si un niveau est sélectionné, on filtre aussi la liste des classes
    if niveau:
        classes = classes.filter(niveau=niveau)

    # -----------------------------
    # STATISTIQUES
    # -----------------------------

    # Une moyenne trimestrielle égale à 0 correspond à un élève non
    # classé : elle ne doit donc ni être retenue comme "plus faible
    # moyenne", ni entrer dans le calcul de la moyenne de la classe.
    resultats_avec_moyenne = resultats.exclude(
        moyenne_trimestrielle=0
    )

    statistiques = resultats.aggregate(

        nombre_eleves=Count("id"),

    )

    statistiques.update(

        resultats_avec_moyenne.aggregate(

            moyenne_classe=Avg(
                "moyenne_trimestrielle"
            ),

            meilleure_moyenne=Max(
                "moyenne_trimestrielle"
            ),

            plus_faible_moyenne=Min(
                "moyenne_trimestrielle"
            ),

        )

    )

    # -----------------------------
    # CLASSEMENT + PERFORMANCE
    # -----------------------------

    classement_queryset = resultats.order_by(
        "-moyenne_trimestrielle",
        "eleve__nom",
        "eleve__prenoms",
    )

    # On prépare le classement complet avant la pagination afin que le
    # rang reste correct sur toutes les pages.
    classement_liste = list(classement_queryset)

    for rang, resultat in enumerate(classement_liste, start=1):
        resultat.rang = rang

        moyenne = resultat.moyenne_trimestrielle

        if moyenne is None:
            resultat.performance = "À accompagner"
            resultat.performance_css = "a-accompagner"
        elif moyenne >= 16:
            resultat.performance = "Excellent"
            resultat.performance_css = "excellent"
        elif moyenne >= 10:
            resultat.performance = "Satisfaisant"
            resultat.performance_css = "satisfaisant"
        else:
            resultat.performance = "À accompagner"
            resultat.performance_css = "a-accompagner"

    total_resultats = len(classement_liste)

    # -----------------------------
    # PAGINATION DU CLASSEMENT (10 élèves par page)
    # -----------------------------

    paginator_classement = Paginator(
        classement_liste,
        10
    )

    numero_page_classement = request.GET.get("page_classement")

    classement = paginator_classement.get_page(
        numero_page_classement
    )

    # Chaîne de requête (filtres actuels) sans le paramètre de page,
    # pour reconstruire les liens de pagination dans le template.
    querydict_classement = request.GET.copy()
    querydict_classement.pop("page_classement", None)
    querystring_classement = querydict_classement.urlencode()

    # -----------------------------
    # TABLEAU STATISTIQUES PAR NIVEAU
    # -----------------------------

    tableau_statistiques_niveaux = calculer_tableau_statistiques_niveaux(
        resultats
    )

    tableau_effectifs_niveaux = calculer_tableau_effectifs_niveaux(
        resultats
    )

    # -----------------------------
    # LIBELLÉS DES FILTRES ACTIFS
    # -----------------------------

    filtres_actifs = []

    annee_obj = None
    if annee_id:
        annee_obj = AnneeScolaire.objects.filter(id=annee_id).first()
        if annee_obj:
            filtres_actifs.append({
                "libelle": "Année scolaire",
                "valeur": annee_obj.nom,
            })

    if niveau:
        filtres_actifs.append({
            "libelle": "Niveau",
            "valeur": niveau,
        })

    classe_obj = None
    if classe_id:
        classe_obj = Classe.objects.filter(id=classe_id).first()
        if classe_obj:
            filtres_actifs.append({
                "libelle": "Classe",
                "valeur": classe_obj.nom,
            })

    trimestre_labels = {
        "T1": "Trimestre 1",
        "T2": "Trimestre 2",
        "T3": "Trimestre 3",
    }

    if trimestre:
        filtres_actifs.append({
            "libelle": "Trimestre",
            "valeur": trimestre_labels.get(trimestre, trimestre),
        })

    if recherche:
        filtres_actifs.append({
            "libelle": "Recherche",
            "valeur": recherche,
        })

    # Période clairement affichée dans le dashboard
    periode_annee = annee_obj.nom if annee_obj else "Toutes les années"
    periode_trimestre = trimestre_labels.get(
        trimestre,
        "Tous les trimestres"
    )

    if annee_obj and trimestre:
        periode_analysee = f"{periode_annee} — {periode_trimestre}"
    elif annee_obj:
        periode_analysee = f"{periode_annee} — tous les trimestres"
    elif trimestre:
        periode_analysee = f"Toutes les années — {periode_trimestre}"
    else:
        periode_analysee = "Toutes les années — tous les trimestres"

    # -----------------------------
    # ZONE D'ALERTE ET D'ACCOMPAGNEMENT
    # -----------------------------

    # Élèves dont la moyenne est inférieure à 10, du plus faible au plus élevé.
    eleves_sous_10 = list(
        resultats
        .filter(moyenne_trimestrielle__lt=10)
        .order_by("moyenne_trimestrielle", "eleve__nom", "eleve__prenoms")
    )
    
    # Matières ayant les moyennes les plus faibles.
    statistiques_matieres_dashboard = calculer_statistiques_matieres(resultats)
    matieres_faibles = sorted(
        [m for m in statistiques_matieres_dashboard if m["moyenne"] is not None],
        key=lambda m: m["moyenne"]
    )[:5]

    # Classes ayant les plus faibles taux de réussite.
    classes_stats = {}
    for resultat in resultats:
        classe_nom = resultat.classe.nom
        stats_classe = classes_stats.setdefault(
            resultat.classe_id,
            {
                "classe": classe_nom,
                "total": 0,
                "reussis": 0,
            }
        )

        moyenne = resultat.moyenne_trimestrielle
        if moyenne is not None and moyenne > 0:
            stats_classe["total"] += 1
            if moyenne >= SEUIL_REUSSITE:
                stats_classe["reussis"] += 1

    classes_faibles = []
    for stats_classe in classes_stats.values():
        if stats_classe["total"]:
            stats_classe["taux_reussite"] = round(
                stats_classe["reussis"] / stats_classe["total"] * 100,
                2
            )
            classes_faibles.append(stats_classe)

    classes_faibles.sort(
        key=lambda c: (c["taux_reussite"], c["classe"])
    )
    classes_faibles = classes_faibles[:10]

    # Liste prioritaire : les 10 élèves ayant les moyennes les plus faibles.
    eleves_a_accompagner = eleves_sous_10[:10]

    # -----------------------------
    # DONNÉES DES GRAPHIQUES
    # -----------------------------

    donnees_graphiques = construire_donnees_graphiques(
        tableau_statistiques_niveaux
    )

    context = {

        "form": form,

        "annees": annees,

        "classes": classes,

        "niveaux": niveaux_uniques,  # Ajout des niveaux dans le contexte

        "resultats": resultats,

        "classement": classement,

        "total_resultats": total_resultats,

        "querystring_classement": querystring_classement,

        "statistiques": statistiques,

        "annee_selectionnee": annee_id,

        "classe_selectionnee": classe_id,

        "trimestre_selectionne": trimestre,

        "niveau_selectionne": niveau,
        "recherche": recherche,
        "filtres_actifs": filtres_actifs,
        "periode_analysee": periode_analysee,

        "tableau_statistiques_niveaux": tableau_statistiques_niveaux,

        "tableau_effectifs_niveaux": tableau_effectifs_niveaux,

        "donnees_graphiques": donnees_graphiques,

        "eleves_sous_10": eleves_sous_10,
        "matieres_faibles": matieres_faibles,
        "classes_faibles": classes_faibles,
        "eleves_a_accompagner": eleves_a_accompagner,

    }


    return render(
        request,
        "analyse/dashboard.html",
        context
    )


def statistiques_matieres(request):

    annee_id = request.GET.get("annee")
    classe_id = request.GET.get("classe")
    trimestre = request.GET.get("trimestre")
    niveau = request.GET.get("niveau")  # Nouveau paramètre

    annees = AnneeScolaire.objects.all().order_by("-nom")
    classes = Classe.objects.all().order_by("nom")
    
    # Récupération des niveaux uniques existants
    niveaux_uniques = Classe.objects.values_list('niveau', flat=True).distinct().exclude(niveau='').order_by('niveau')

    resultats = Resultat.objects.select_related(
        "eleve",
        "annee_scolaire",
        "classe"
    )

    # Filtres
    if niveau:
        resultats = resultats.filter(
            classe__niveau=niveau
        )

    if annee_id:
        resultats = resultats.filter(
            annee_scolaire_id=annee_id
        )

    if classe_id:
        resultats = resultats.filter(
            classe_id=classe_id
        )

    if trimestre:
        resultats = resultats.filter(
            trimestre=trimestre
        )

    # Si un niveau est sélectionné, on filtre aussi la liste des classes
    if niveau:
        classes = classes.filter(niveau=niveau)

    statistiques_matieres = calculer_statistiques_matieres(resultats)

    # Matières avec des données
    matieres_valides = [
        m for m in statistiques_matieres
        if m["moyenne"] is not None
    ]

    if matieres_valides:

        meilleure_matiere = max(
            matieres_valides,
            key=lambda x: x["moyenne"]
        )

        plus_faible_matiere = min(
            matieres_valides,
            key=lambda x: x["moyenne"]
        )

    else:

        meilleure_matiere = None
        plus_faible_matiere = None

    context = {
        "annees": annees,
        "classes": classes,
        "niveaux": niveaux_uniques,  # Ajout des niveaux

        "statistiques_matieres":
            statistiques_matieres,

        "meilleure_matiere":
            meilleure_matiere,

        "plus_faible_matiere":
            plus_faible_matiere,

        "annee_selectionnee": annee_id,
        "classe_selectionnee": classe_id,
        "trimestre_selectionne": trimestre,
        "niveau_selectionne": niveau,  # Ajout du niveau sélectionné
    }

    return render(
        request,
        "analyse/statistiques_matieres.html",
        context
    )


# =====================================================================
# EXPORT EXCEL DE TOUTES LES STATISTIQUES DU DASHBOARD
# =====================================================================

POLICE_ENTETE = Font(bold=True, name="Arial")
ALIGNEMENT_CENTRE = Alignment(horizontal="center", vertical="center")


def _ecrire_entetes(feuille, entetes):
    """Écrit une ligne d'en-têtes stylée (gras, sans couleur) en haut
    d'une feuille."""

    for colonne, texte in enumerate(entetes, start=1):

        cellule = feuille.cell(row=1, column=colonne, value=texte)
        cellule.font = POLICE_ENTETE
        cellule.alignment = ALIGNEMENT_CENTRE

    feuille.freeze_panes = "A2"


def _ajuster_largeurs_colonnes(feuille):
    """Ajuste automatiquement la largeur des colonnes selon leur contenu."""

    for colonne in feuille.columns:

        longueur_max = max(
            (len(str(cellule.value)) for cellule in colonne if cellule.value is not None),
            default=10
        )

        lettre = get_column_letter(colonne[0].column)
        feuille.column_dimensions[lettre].width = min(longueur_max + 4, 40)


# ---- Styles reproduisant la structure des tableaux "par niveau" du
#      dashboard, sans couleurs (gras + bordures uniquement) ----

POLICE_ENTETE_TABLEAU = Font(bold=True, name="Arial")

BORDURE_FINE = Border(
    left=Side(style="thin", color="333333"),
    right=Side(style="thin", color="333333"),
    top=Side(style="thin", color="333333"),
    bottom=Side(style="thin", color="333333"),
)


def _cellule_entete_sombre(feuille, ligne, colonne, texte):
    """Écrit une cellule d'en-tête au style des tableaux "par niveau"
    du dashboard (gras, centré, bordé, sans couleur)."""

    cellule = feuille.cell(row=ligne, column=colonne, value=texte)
    cellule.font = POLICE_ENTETE_TABLEAU
    cellule.alignment = ALIGNEMENT_CENTRE
    cellule.border = BORDURE_FINE

    return cellule


def _remplir_ligne_donnees(feuille, ligne_excel, valeurs, est_total, est_total_general,
                            colonnes_pourcentage=()):
    """Écrit une ligne de données, avec les lignes de total en gras
    (comme les classes CSS total-row / total-general du dashboard),
    mais sans couleur de fond."""

    for indice_colonne, valeur in enumerate(valeurs, start=1):

        cellule = feuille.cell(row=ligne_excel, column=indice_colonne, value=valeur)
        cellule.border = BORDURE_FINE
        cellule.alignment = ALIGNEMENT_CENTRE

        if indice_colonne in colonnes_pourcentage and isinstance(valeur, (int, float)):
            cellule.value = valeur / 100
            cellule.number_format = "0.00%"

        if est_total_general or est_total or indice_colonne == 1:
            cellule.font = Font(bold=True, name="Arial")


def _construire_feuille_statistiques_niveaux(feuille, lignes):
    """Construit la feuille "Statistiques par niveau" en reproduisant

    exactement le tableau du dashboard : en-têtes groupés sur deux
    lignes, ordre de colonnes F/G/T, pourcentages et mise en
    évidence des lignes de total."""

    groupes = [
        ("Effectif classé", ["F", "G", "T"]),
        ("Effectif non classé", ["F", "G", "T"]),
        ("MOY TRIMES >= 10", ["F", "G", "T", "%"]),
        ("08.50 <= MOY TRIMES < 10", ["F", "G", "T", "%"]),
        ("MOY TRIMES < 08.50", ["F", "G", "T", "%"]),
    ]

    _cellule_entete_sombre(feuille, 1, 1, "Niveaux")
    feuille.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

    _cellule_entete_sombre(feuille, 1, 2, "Nombre de classes")
    feuille.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)

    colonne = 3
    colonnes_pourcentage = []

    for titre_groupe, sous_colonnes in groupes:

        largeur = len(sous_colonnes)

        _cellule_entete_sombre(feuille, 1, colonne, titre_groupe)
        feuille.merge_cells(
            start_row=1, start_column=colonne,
            end_row=1, end_column=colonne + largeur - 1
        )

        for decalage, sous_titre in enumerate(sous_colonnes):
            _cellule_entete_sombre(feuille, 2, colonne + decalage, sous_titre)
            if sous_titre == "%":
                colonnes_pourcentage.append(colonne + decalage)

        colonne += largeur

    _cellule_entete_sombre(feuille, 1, colonne, "Moyenne par niveau")
    feuille.merge_cells(start_row=1, start_column=colonne, end_row=2, end_column=colonne)

    nombre_colonnes = colonne

    ligne_excel = 3

    for ligne in lignes:

        valeurs = [
            ligne["niveau"],
            ligne["nombre_classes"],
            ligne["classe_f"], ligne["classe_g"], ligne["classe_t"],
            ligne["non_classe_f"], ligne["non_classe_g"], ligne["non_classe_t"],
            ligne["sup10_f"], ligne["sup10_g"], ligne["sup10_t"], ligne["sup10_pct"],
            ligne["moy_f"], ligne["moy_g"], ligne["moy_t"], ligne["moy_pct"],
            ligne["inf_f"], ligne["inf_g"], ligne["inf_t"], ligne["inf_pct"],
            ligne["moyenne_niveau"] if ligne["moyenne_niveau"] is not None else "-",
        ]

        _remplir_ligne_donnees(
            feuille,
            ligne_excel,
            valeurs,
            est_total=ligne.get("is_total", False),
            est_total_general=(ligne["niveau"] == "Total général"),
            colonnes_pourcentage=colonnes_pourcentage,
        )

        ligne_excel += 1

    feuille.freeze_panes = "C3"

    feuille.column_dimensions["A"].width = 22
    for indice_colonne in range(2, nombre_colonnes + 1):
        feuille.column_dimensions[get_column_letter(indice_colonne)].width = 12


def _construire_feuille_effectifs_niveaux(feuille, lignes):
    """Construit la feuille "Effectifs" en reproduisant exactement le
    tableau du dashboard : en-têtes groupés sur deux lignes, ordre
    de colonnes G/F/Total et mise en évidence des lignes de total."""

    groupes = [
        ("Effectif", ["G", "F", "Total"]),
        ("Affectés", ["G", "F", "Total"]),
        ("Redoublants", ["G", "F", "Total"]),
    ]

    _cellule_entete_sombre(feuille, 1, 1, "Niveau")
    feuille.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

    colonne = 2

    for titre_groupe, sous_colonnes in groupes:

        largeur = len(sous_colonnes)

        _cellule_entete_sombre(feuille, 1, colonne, titre_groupe)
        feuille.merge_cells(
            start_row=1, start_column=colonne,
            end_row=1, end_column=colonne + largeur - 1
        )

        for decalage, sous_titre in enumerate(sous_colonnes):
            _cellule_entete_sombre(feuille, 2, colonne + decalage, sous_titre)

        colonne += largeur

    nombre_colonnes = colonne - 1

    ligne_excel = 3

    for ligne in lignes:

        valeurs = [
            ligne["niveau"],
            ligne["effectif_g"], ligne["effectif_f"], ligne["effectif_t"],
            ligne["affectes_g"], ligne["affectes_f"], ligne["affectes_t"],
            ligne["redoublants_g"], ligne["redoublants_f"], ligne["redoublants_t"],
        ]

        _remplir_ligne_donnees(
            feuille,
            ligne_excel,
            valeurs,
            est_total=ligne.get("is_total", False),
            est_total_general=(ligne["niveau"] == "Total Général"),
        )

        ligne_excel += 1

    feuille.freeze_panes = "B3"

    feuille.column_dimensions["A"].width = 22
    for indice_colonne in range(2, nombre_colonnes + 1):
        feuille.column_dimensions[get_column_letter(indice_colonne)].width = 12


def exporter_eleves_sous_10_excel(request):
    """Exporte, dans un classeur Excel, la liste complète des élèves
    dont la moyenne trimestrielle est inférieure à 10, du plus faible
    au plus élevé, en respectant les filtres actuellement appliqués
    sur le dashboard (année, niveau, classe, trimestre, recherche)."""

    # -----------------------------
    # FILTRES (identiques à la vue dashboard)
    # -----------------------------

    annee_id = request.GET.get("annee")
    classe_id = request.GET.get("classe")
    trimestre = request.GET.get("trimestre")
    niveau = request.GET.get("niveau")
    recherche = request.GET.get("recherche", "").strip()

    resultats = Resultat.objects.select_related(
        "eleve",
        "annee_scolaire",
        "classe"
    )

    if niveau:
        resultats = resultats.filter(classe__niveau=niveau)

    if annee_id:
        resultats = resultats.filter(annee_scolaire_id=annee_id)

    if classe_id:
        resultats = resultats.filter(classe_id=classe_id)

    if trimestre:
        resultats = resultats.filter(trimestre=trimestre)

    if recherche:
        resultats = resultats.filter(
            Q(eleve__nom__icontains=recherche)
            | Q(eleve__prenoms__icontains=recherche)
            | Q(eleve__matricule__icontains=recherche)
        )

    eleves_sous_10 = resultats.filter(
        moyenne_trimestrielle__lt=10
    ).order_by("moyenne_trimestrielle", "eleve__nom", "eleve__prenoms")

    # -----------------------------
    # CONSTRUCTION DU CLASSEUR EXCEL
    # -----------------------------

    classeur = Workbook()

    feuille = classeur.active
    feuille.title = "Eleves sous 10"

    _ecrire_entetes(
        feuille,
        ["Rang", "Matricule", "Nom", "Prénoms", "Classe", "Niveau", "Moyenne"]
    )

    ligne_excel = 2

    for rang, resultat in enumerate(eleves_sous_10, start=1):

        feuille.cell(row=ligne_excel, column=1, value=rang)
        feuille.cell(row=ligne_excel, column=2, value=resultat.eleve.matricule)
        feuille.cell(row=ligne_excel, column=3, value=resultat.eleve.nom)
        feuille.cell(row=ligne_excel, column=4, value=resultat.eleve.prenoms)
        feuille.cell(row=ligne_excel, column=5, value=resultat.classe.nom)
        feuille.cell(row=ligne_excel, column=6, value=resultat.classe.niveau)
        feuille.cell(
            row=ligne_excel,
            column=7,
            value=float(resultat.moyenne_trimestrielle)
        )

        ligne_excel += 1

    _ajuster_largeurs_colonnes(feuille)

    # -----------------------------
    # RÉPONSE HTTP (téléchargement du fichier)
    # -----------------------------

    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)

    horodatage = datetime.now().strftime("%Y%m%d_%H%M")
    nom_fichier = f"eleves_sous_10_{horodatage}.xlsx"

    reponse = HttpResponse(
        tampon.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        )
    )
    reponse["Content-Disposition"] = f'attachment; filename="{nom_fichier}"'

    return reponse


def exporter_statistiques_excel(request):
    """Exporte, dans un classeur Excel à plusieurs feuilles, l'ensemble
    des statistiques affichées sur le dashboard (résumé, classement
    complet des élèves, tableau par niveau, tableau des effectifs),
    en respectant les filtres actuellement appliqués (année, niveau,
    classe, trimestre)."""

    # -----------------------------
    # FILTRES (identiques à la vue dashboard)
    # -----------------------------

    annee_id = request.GET.get("annee")
    classe_id = request.GET.get("classe")
    trimestre = request.GET.get("trimestre")
    niveau = request.GET.get("niveau")
    recherche = request.GET.get("recherche", "").strip()

    resultats = Resultat.objects.select_related(
        "eleve",
        "annee_scolaire",
        "classe"
    )

    if niveau:
        resultats = resultats.filter(classe__niveau=niveau)

    if annee_id:
        resultats = resultats.filter(annee_scolaire_id=annee_id)

    if classe_id:
        resultats = resultats.filter(classe_id=classe_id)

    if trimestre:
        resultats = resultats.filter(trimestre=trimestre)

    if recherche:
        resultats = resultats.filter(
            Q(eleve__nom__icontains=recherche)
            | Q(eleve__prenoms__icontains=recherche)
            | Q(eleve__matricule__icontains=recherche)
        )

    # -----------------------------
    # STATISTIQUES (mêmes calculs que la vue dashboard)
    # -----------------------------

    resultats_avec_moyenne = resultats.exclude(
        moyenne_trimestrielle=0
    )

    statistiques = resultats.aggregate(
        nombre_eleves=Count("id"),
    )

    statistiques.update(
        resultats_avec_moyenne.aggregate(
            moyenne_classe=Avg("moyenne_trimestrielle"),
            meilleure_moyenne=Max("moyenne_trimestrielle"),
            plus_faible_moyenne=Min("moyenne_trimestrielle"),
        )
    )

    classement = resultats.order_by(
        "-moyenne_trimestrielle"
    )

    tableau_statistiques_niveaux = calculer_tableau_statistiques_niveaux(
        resultats
    )

    tableau_effectifs_niveaux = calculer_tableau_effectifs_niveaux(
        resultats
    )

    # -----------------------------
    # LIBELLÉS DES FILTRES APPLIQUÉS (pour le rappel dans le fichier)
    # -----------------------------

    annee_nom = "Toutes les années"
    if annee_id:
        annee_obj = AnneeScolaire.objects.filter(id=annee_id).first()
        annee_nom = annee_obj.nom if annee_obj else annee_id

    classe_nom = "Toutes les classes"
    if classe_id:
        classe_obj = Classe.objects.filter(id=classe_id).first()
        classe_nom = classe_obj.nom if classe_obj else classe_id

    niveau_nom = niveau or "Tous les niveaux"
    trimestre_nom = trimestre or "Tous les trimestres"

    # -----------------------------
    # CONSTRUCTION DU CLASSEUR EXCEL
    # -----------------------------

    classeur = Workbook()

    # ---- Feuille "Résumé" ----

    feuille_resume = classeur.active
    feuille_resume.title = "Résumé"

    _ecrire_entetes(feuille_resume, ["Indicateur", "Valeur"])

    moyenne_classe = statistiques.get("moyenne_classe")
    meilleure_moyenne = statistiques.get("meilleure_moyenne")
    plus_faible_moyenne = statistiques.get("plus_faible_moyenne")

    lignes_resume = [
        ("Date d'export", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Année scolaire", annee_nom),
        ("Niveau", niveau_nom),
        ("Classe", classe_nom),
        ("Trimestre", trimestre_nom),
        ("Nombre d'élèves", statistiques.get("nombre_eleves") or 0),
        (
            "Moyenne générale",
            round(float(moyenne_classe), 2)
            if moyenne_classe is not None else "-"
        ),
        (
            "Meilleure moyenne",
            round(float(meilleure_moyenne), 2)
            if meilleure_moyenne is not None else "-"
        ),
        (
            "Plus faible moyenne",
            round(float(plus_faible_moyenne), 2)
            if plus_faible_moyenne is not None else "-"
        ),
    ]

    for indice, (libelle, valeur) in enumerate(lignes_resume, start=2):
        feuille_resume.cell(row=indice, column=1, value=libelle).font = Font(
            bold=True, name="Arial"
        )
        feuille_resume.cell(row=indice, column=2, value=valeur)

    _ajuster_largeurs_colonnes(feuille_resume)

    # ---- Feuille "Classement" ----

    feuille_classement = classeur.create_sheet("Classement")

    _ecrire_entetes(
        feuille_classement,
        ["Rang", "Matricule", "Nom", "Prénoms", "Classe", "Niveau", "Moyenne"]
    )

    rang = 0

    for indice, resultat in enumerate(classement, start=2):

        moyenne = resultat.moyenne_trimestrielle
        est_classe = moyenne not in (None, 0)

        if est_classe:
            rang += 1

        feuille_classement.cell(row=indice, column=1, value=rang if est_classe else "-")
        feuille_classement.cell(row=indice, column=2, value=resultat.eleve.matricule)
        feuille_classement.cell(row=indice, column=3, value=resultat.eleve.nom)
        feuille_classement.cell(row=indice, column=4, value=resultat.eleve.prenoms)
        feuille_classement.cell(row=indice, column=5, value=resultat.classe.nom)
        feuille_classement.cell(row=indice, column=6, value=resultat.classe.niveau)
        feuille_classement.cell(
            row=indice,
            column=7,
            value=float(moyenne) if est_classe else None
        )

    _ajuster_largeurs_colonnes(feuille_classement)

    # ---- Feuille "Statistiques par niveau" (identique au dashboard) ----

    feuille_niveaux = classeur.create_sheet("Statistiques par niveau")

    _construire_feuille_statistiques_niveaux(
        feuille_niveaux,
        tableau_statistiques_niveaux
    )

    # ---- Feuille "Effectifs" (identique au dashboard) ----

    feuille_effectifs = classeur.create_sheet("Effectifs")

    _construire_feuille_effectifs_niveaux(
        feuille_effectifs,
        tableau_effectifs_niveaux
    )

    # -----------------------------
    # RÉPONSE HTTP (téléchargement du fichier)
    # -----------------------------

    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)

    horodatage = datetime.now().strftime("%Y%m%d_%H%M")
    nom_fichier = f"statistiques_{horodatage}.xlsx"

    reponse = HttpResponse(
        tampon.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        )
    )
    reponse["Content-Disposition"] = f'attachment; filename="{nom_fichier}"'

    return reponse